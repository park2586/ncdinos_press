"""
보도자료 초안을 Word(.docx) 파일로 저장합니다.
"""
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")


def save_docx(title: str, content: str, category: str) -> str:
    """초안을 Word 파일로 저장하고 경로를 반환합니다."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    doc = Document()

    # 페이지 여백 설정
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)

    # ── 헤더: NC 다이노스 보도자료 ──
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_para.add_run("NC 다이노스  보도자료")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)

    # 카테고리 배지
    cat_para = doc.add_paragraph()
    cat_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = cat_para.add_run(f"[ {category} ]")
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    # 구분선
    doc.add_paragraph("─" * 55)

    # 본문 파싱 및 삽입
    lines = content.strip().splitlines()
    is_first_line = True

    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue

        para = doc.add_paragraph()

        # 제목 (첫 줄)
        if is_first_line:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(stripped)
            run.bold = True
            run.font.size = Pt(14)
            is_first_line = False

        # 소제목 (- 로 시작)
        elif stripped.startswith("- ") and len(stripped) < 60:
            run = para.add_run(stripped)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            para.paragraph_format.left_indent = Cm(0.5)

        # (끝) 마커
        elif stripped == "(끝)":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run("(끝)")
            run.bold = True
            run.font.size = Pt(10)

        # [첨부파일 설명] 헤더
        elif stripped.startswith("[첨부") or stripped.startswith("[붙임"):
            run = para.add_run(stripped)
            run.bold = True
            run.font.size = Pt(10)

        # 일반 본문
        else:
            run = para.add_run(stripped)
            run.font.size = Pt(10.5)
            para.paragraph_format.space_after = Pt(4)

    # 날짜 푸터
    doc.add_paragraph("")
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer_para.add_run(f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    # 파일 저장
    safe_title = "".join(c for c in title[:30] if c not in r'\/:*?"<>|')
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"NC_보도자료_{safe_title}_{timestamp}.docx"
    filepath = os.path.join(OUTPUT_DIR, filename)
    doc.save(filepath)

    return filepath


if __name__ == "__main__":
    sample = """NC 다이노스, 2025 팬페스티벌 개최

NC 다이노스가 오는 4월 20일(일) 오후 2시 창원NC파크에서 2025 팬페스티벌을 개최한다.

이번 행사에는 팬사인회, 선수 토크쇼, 유니폼 증정 이벤트 등 다양한 프로그램이 준비되어 있다.

임선남 단장은 "팬 여러분께 특별한 경험을 드리기 위해 다양한 프로그램을 준비했다"라고 말했다.

(끝)

[첨부파일 설명]

- 사진 1: 2025 팬페스티벌 포스터"""

    path = save_docx("팬페스티벌", sample, "마케팅")
    print(f"저장됨: {path}")
